"""
Routes admin - Sauvegarde des Données
Export manuel JSON (ZIP) et Excel multi-onglets.
Sauvegarde automatique vers Google Drive + notification email.
La clé Google est stockée en base : aucun recodage nécessaire à l'activation.
"""

import asyncio
import io
import json
import zipfile
import logging
from datetime import datetime, timezone
from typing import Any

from bson import ObjectId
from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseUpload

from middleware.auth import RoleChecker

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/backup", tags=["Sauvegarde"])
admin_only = RoleChecker(["admin"])

db = None

# Scope Google Drive : création et gestion des fichiers déposés par l'application
DRIVE_SCOPES = ["https://www.googleapis.com/auth/drive.file"]

# Liste ordonnée de toutes les collections à sauvegarder
COLLECTIONS_TO_BACKUP = [
    "users",
    "profiles",
    "portfolio",
    "user_subscriptions",
    "subscription_plans",
    "subscriptions",
    "projects",
    "project_applications",
    "project_rooms",
    "project_messages",
    "project_notes",
    "announcements",
    "alerts",
    "partners",
    "messages",
    "notifications",
    "contacts",
    "invoices",
    "user_billing",
    "admin_logs",
    "login_attempts",
    "blocked_ips",
    "brute_force_config",
    "password_resets",
]


def set_database(database):
    """Injecte la connexion à la base de données"""
    global db
    db = database


# ============================================
# UTILITAIRES DE SÉRIALISATION
# ============================================

def _serialize_value(value: Any) -> Any:
    """Convertit les types MongoDB (ObjectId, datetime) en types JSON sérialisables"""
    if isinstance(value, ObjectId):
        return str(value)
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, dict):
        return {k: _serialize_value(v) for k, v in value.items()}
    if isinstance(value, list):
        return [_serialize_value(item) for item in value]
    return value


def _serialize_document(doc: dict) -> dict:
    """Sérialise un document MongoDB complet"""
    return {key: _serialize_value(value) for key, value in doc.items()}


# ============================================
# MODÈLES
# ============================================

class BackupConfig(BaseModel):
    notification_email: str = ""
    google_drive_enabled: bool = False
    google_drive_credentials_json: str = ""  # Contenu JSON de la clé de service Google
    google_drive_folder_id: str = ""
    auto_backup_enabled: bool = False
    auto_backup_frequency: str = "weekly"  # daily | weekly | monthly


class TriggerResult(BaseModel):
    success: bool
    filename: str
    drive_uploaded: bool
    drive_url: str = ""
    email_sent: bool
    error: str = ""
    triggered_at: str


# ============================================
# UTILITAIRES INTERNES
# ============================================

async def _generate_zip_buffer() -> io.BytesIO:
    """Génère le ZIP de backup de toutes les collections en mémoire."""
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, mode="w", compression=zipfile.ZIP_DEFLATED) as zf:
        for collection_name in COLLECTIONS_TO_BACKUP:
            try:
                cursor = db[collection_name].find({})
                documents = await cursor.to_list(length=None)
                serialized = [_serialize_document(doc) for doc in documents]
                json_content = json.dumps(serialized, ensure_ascii=False, indent=2)
                zf.writestr(f"{collection_name}.json", json_content)
            except Exception as e:
                logger.error("Erreur export JSON collection %s : %s", collection_name, e)
                zf.writestr(f"{collection_name}.json", json.dumps({"error": str(e)}))
    zip_buffer.seek(0)
    return zip_buffer


def _drive_upload_sync(credentials_json: str, folder_id: str, filename: str, file_bytes: bytes) -> str:
    """
    Upload synchrone d'un fichier vers Google Drive.
    Conçu pour être exécuté dans un thread (non-async).
    Retourne l'URL de partage du fichier.
    """
    credentials_info = json.loads(credentials_json)
    credentials = service_account.Credentials.from_service_account_info(
        credentials_info,
        scopes=DRIVE_SCOPES
    )
    service = build("drive", "v3", credentials=credentials, cache_discovery=False)

    file_metadata = {"name": filename}
    if folder_id:
        file_metadata["parents"] = [folder_id]

    media = MediaIoBaseUpload(
        io.BytesIO(file_bytes),
        mimetype="application/zip",
        resumable=False
    )

    created = service.files().create(
        body=file_metadata,
        media_body=media,
        fields="id,webViewLink"
    ).execute()

    return created.get("webViewLink", "")


# ============================================
# ENDPOINTS
# ============================================

@router.get("/stats", dependencies=[Depends(admin_only)])
async def get_backup_stats():
    """Retourne le comptage de documents par collection"""
    stats = []
    total_documents = 0

    for collection_name in COLLECTIONS_TO_BACKUP:
        try:
            count = await db[collection_name].count_documents({})
            stats.append({"collection": collection_name, "count": count})
            total_documents += count
        except Exception as e:
            logger.warning("Impossible de compter la collection %s : %s", collection_name, e)
            stats.append({"collection": collection_name, "count": 0})

    return {
        "collections": stats,
        "total_documents": total_documents,
        "total_collections": len(COLLECTIONS_TO_BACKUP),
        "generated_at": datetime.now(timezone.utc).isoformat(),
    }


@router.get("/export/json", dependencies=[Depends(admin_only)])
async def export_json():
    """
    Exporte toutes les collections en JSON.
    Retourne un fichier ZIP contenant un fichier JSON par collection.
    """
    timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%d_%H%M%S")
    zip_filename = f"backup_syndicat_{timestamp}.zip"

    zip_buffer = await _generate_zip_buffer()
    logger.info("Export JSON généré : %s", zip_filename)

    return StreamingResponse(
        zip_buffer,
        media_type="application/zip",
        headers={"Content-Disposition": f"attachment; filename={zip_filename}"},
    )


@router.get("/export/excel", dependencies=[Depends(admin_only)])
async def export_excel():
    """
    Exporte toutes les collections en Excel multi-onglets.
    Un onglet par collection avec en-têtes stylisés.
    """
    timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%d_%H%M%S")
    excel_filename = f"backup_syndicat_{timestamp}.xlsx"

    wb = Workbook()
    wb.remove(wb.active)  # Supprimer la feuille vide créée par défaut

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4068", end_color="1F4068", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    for collection_name in COLLECTIONS_TO_BACKUP:
        # Nom de feuille limité à 31 caractères (contrainte Excel)
        sheet_name = collection_name[:31]
        ws = wb.create_sheet(title=sheet_name)

        try:
            cursor = db[collection_name].find({})
            documents = await cursor.to_list(length=None)

            if not documents:
                ws.append(["(collection vide)"])
                continue

            # En-têtes : union ordonnée de toutes les clés présentes dans les documents
            all_keys = list(dict.fromkeys(
                key for doc in documents for key in _serialize_document(doc).keys()
            ))

            for col_idx, key in enumerate(all_keys, start=1):
                cell = ws.cell(row=1, column=col_idx, value=key)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment

            for row_idx, doc in enumerate(documents, start=2):
                serialized = _serialize_document(doc)
                for col_idx, key in enumerate(all_keys, start=1):
                    value = serialized.get(key, "")
                    # Les dict/list sont convertis en chaîne JSON pour Excel
                    if isinstance(value, (dict, list)):
                        value = json.dumps(value, ensure_ascii=False)
                    ws.cell(row=row_idx, column=col_idx, value=value)

            # Ajustement automatique de la largeur des colonnes
            for col in ws.columns:
                max_length = max(
                    (len(str(cell.value)) if cell.value else 0 for cell in col),
                    default=0,
                )
                ws.column_dimensions[col[0].column_letter].width = min(max_length + 4, 50)

        except Exception as e:
            logger.error("Erreur export Excel collection %s : %s", collection_name, e)
            ws.append(["Erreur lors de l'export", str(e)])

    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)

    logger.info("Export Excel généré : %s", excel_filename)

    return StreamingResponse(
        excel_buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={excel_filename}"},
    )


@router.post("/trigger", dependencies=[Depends(admin_only)])
async def trigger_backup():
    """
    Déclenche une sauvegarde complète :
    1. Génère le ZIP de toutes les collections
    2. Uploade vers Google Drive (si clé configurée)
    3. Envoie un email de notification (si email configuré)
    Retourne le résultat détaillé de chaque étape.
    """
    config = await db.backup_config.find_one({}, {"_id": 0}) or {}

    timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%d_%H%M%S")
    filename = f"backup_syndicat_{timestamp}.zip"
    triggered_at = datetime.now(timezone.utc).isoformat()

    drive_uploaded = False
    drive_url = ""
    email_sent = False
    error = ""

    # Étape 1 — Générer le ZIP
    try:
        zip_buffer = await _generate_zip_buffer()
        file_bytes = zip_buffer.read()
        logger.info("ZIP de backup généré : %s (%d octets)", filename, len(file_bytes))
    except Exception as e:
        logger.error("Erreur génération ZIP : %s", e)
        return TriggerResult(
            success=False,
            filename=filename,
            drive_uploaded=False,
            email_sent=False,
            error=f"Erreur génération ZIP : {e}",
            triggered_at=triggered_at,
        )

    # Étape 2 — Upload Google Drive (uniquement si clé présente)
    credentials_json = config.get("google_drive_credentials_json", "")
    folder_id = config.get("google_drive_folder_id", "")

    if credentials_json:
        try:
            loop = asyncio.get_event_loop()
            drive_url = await loop.run_in_executor(
                None,
                _drive_upload_sync,
                credentials_json,
                folder_id,
                filename,
                file_bytes,
            )
            drive_uploaded = True
            logger.info("Backup uploadé sur Google Drive : %s", drive_url)
        except Exception as e:
            logger.error("Erreur upload Google Drive : %s", e)
            error = f"Erreur Google Drive : {e}"
    else:
        logger.info("Google Drive non configuré — upload ignoré")

    # Étape 3 — Email de notification (si email renseigné)
    notification_email = config.get("notification_email", "")
    if notification_email:
        from services.email_service import send_backup_notification_email
        email_sent = send_backup_notification_email(
            to_email=notification_email,
            filename=filename,
            drive_url=drive_url,
            drive_enabled=bool(credentials_json),
            error_message=error,
        )

    # Sauvegarde de l'historique en base
    await db.backup_history.insert_one({
        "filename": filename,
        "drive_uploaded": drive_uploaded,
        "drive_url": drive_url,
        "email_sent": email_sent,
        "error": error,
        "triggered_at": triggered_at,
    })

    return TriggerResult(
        success=True,
        filename=filename,
        drive_uploaded=drive_uploaded,
        drive_url=drive_url,
        email_sent=email_sent,
        error=error,
        triggered_at=triggered_at,
    )


@router.get("/config", dependencies=[Depends(admin_only)])
async def get_backup_config():
    """Récupère la configuration de sauvegarde (Google Drive, email, planification)"""
    config = await db.backup_config.find_one({}, {"_id": 0})
    if not config:
        return {
            "notification_email": "",
            "google_drive_enabled": False,
            "google_drive_credentials_json": "",
            "google_drive_folder_id": "",
            "auto_backup_enabled": False,
            "auto_backup_frequency": "weekly",
        }
    return config


@router.post("/config", dependencies=[Depends(admin_only)])
async def save_backup_config(config: BackupConfig):
    """
    Sauvegarde la configuration de sauvegarde automatique.
    La clé Google Drive est stockée en base : l'activation ne nécessite aucun recodage.
    """
    config_data = config.dict()
    config_data["updated_at"] = datetime.now(timezone.utc).isoformat()

    await db.backup_config.replace_one({}, config_data, upsert=True)
    logger.info("Configuration de sauvegarde mise à jour")

    return {"message": "Configuration sauvegardée avec succès"}
